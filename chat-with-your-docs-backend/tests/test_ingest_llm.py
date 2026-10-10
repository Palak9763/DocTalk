"""Unit tests for rag/ingest.py and rag/llm.py."""

import json
from pathlib import Path
import sys
import tempfile

from docx import Document
import numpy as np

# Ensure project root is in sys.path
ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from rag import config
from rag.ingest import Chunk, chunk_text, load_chunks, load_document
from rag.llm import chat, chat_json, chat_stream, embed, parse_json, reset_embedding_cache
from tests.fakes import install, make_pdf


def test_chunk_text() -> None:
    """Verify sliding-window chunking, exact overlap, and error conditions."""
    words = [f"w{i}" for i in range(250)]
    text = " ".join(words)
    chunk_w, overlap = 50, 15

    chunks = chunk_text(text, words=chunk_w, overlap=overlap)
    assert len(chunks) > 1

    # Verify consecutive chunks share exactly `overlap` words
    for i in range(len(chunks) - 1):
        c1_words = chunks[i].split()
        c2_words = chunks[i + 1].split()
        assert len(c1_words) == chunk_w
        assert c1_words[-overlap:] == c2_words[:overlap]

    # Verify last chunk reaches the end of the text
    last_words = chunks[-1].split()
    assert last_words[-1] == "w249"

    # Verify overlap >= words raises ValueError
    try:
        chunk_text("test text", words=20, overlap=20)
        assert False, "Should have raised ValueError when overlap == words"
    except ValueError:
        pass

    try:
        chunk_text("test text", words=20, overlap=25)
        assert False, "Should have raised ValueError when overlap > words"
    except ValueError:
        pass


def test_markdown_ingest(tmp_path: Path) -> None:
    """Verify markdown sections, unique stable IDs, and body prefix removal."""
    md_content = """Pre-heading introductory remarks explaining the document.

# Section Alpha
Paragraph discussing topic alpha with sufficient detail.

## Section Beta
Detailed breakdown of topic beta.
"""
    file_path = tmp_path / "notes_sample.md"
    file_path.write_text(md_content, encoding="utf-8")

    chunks = load_document(file_path, words=40, overlap=10)
    assert len(chunks) >= 3

    # Check sections detected
    sections = [c.section for c in chunks]
    assert "Introduction" in sections
    assert "Section Alpha" in sections
    assert "Section Beta" in sections

    # Check unique and stable IDs
    ids = [c.id for c in chunks]
    assert len(set(ids)) == len(ids)
    chunks_second_run = load_document(file_path, words=40, overlap=10)
    assert [c.id for c in chunks_second_run] == ids

    # Check header and body properties
    for c in chunks:
        header = f"{file_path.name} › {c.section}\n"
        assert c.text.startswith(header)
        assert c.body == c.text[len(header) :]
        assert not c.body.startswith(header)
        assert c.page is None


def test_docx_and_pdf(tmp_path: Path) -> None:
    """Verify DOCX headings and tables, PDF page tracking, and error cases."""
    # 1. Test DOCX
    doc = Document()
    doc.add_heading("Architecture Overview", level=1)
    doc.add_paragraph("Overview paragraph detailing the pipeline.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Component"
    table.cell(0, 1).text = "Latency"
    table.cell(1, 0).text = "Dense"
    table.cell(1, 1).text = "5ms"

    docx_path = tmp_path / "spec.docx"
    doc.save(str(docx_path))

    docx_chunks = load_document(docx_path, words=50, overlap=10)
    assert len(docx_chunks) >= 1
    assert docx_chunks[0].section == "Architecture Overview"
    assert "Component | Latency" in docx_chunks[0].body
    assert "Dense | 5ms" in docx_chunks[0].body
    assert docx_chunks[0].page is None

    # 2. Test PDF with multiple pages
    pdf_path = tmp_path / "doc.pdf"
    make_pdf(pdf_path, ["First page notes.", "Second page content."])
    pdf_chunks = load_document(pdf_path, words=50, overlap=10)
    assert len(pdf_chunks) == 2
    assert pdf_chunks[0].section == "Page 1"
    assert pdf_chunks[0].page == 1
    assert "First page notes" in pdf_chunks[0].body
    assert pdf_chunks[1].section == "Page 2"
    assert pdf_chunks[1].page == 2
    assert "Second page content" in pdf_chunks[1].body

    # 3. Test scanned PDF without text layer
    empty_pdf = tmp_path / "scanned.pdf"
    make_pdf(empty_pdf, ["", ""])
    try:
        load_document(empty_pdf)
        assert False, "Should have raised ValueError on scanned PDF"
    except ValueError as exc:
        assert "scanned PDFs are not supported" in str(exc)

    # 4. Test unsupported file suffix
    bad_file = tmp_path / "data.csv"
    bad_file.write_text("a,b,c")
    try:
        load_document(bad_file)
        assert False, "Should have raised ValueError on unsupported suffix"
    except ValueError as exc:
        assert "Unsupported file format" in str(exc)


def test_load_chunks_sample_corpus() -> None:
    """Verify sample corpus loads into > 15 chunks with non-empty bodies."""
    chunks = load_chunks(config.NOTES_DIR)
    assert len(chunks) > 15, f"Expected > 15 chunks, got {len(chunks)}"
    for c in chunks:
        assert len(c.body.strip()) > 0
        assert c.source.endswith(".md")


def test_parse_json() -> None:
    """Verify parse_json handles plain, fenced, chatter, and invalid formats."""
    expected = {"score": 0.95, "label": "relevant"}

    # Plain JSON
    assert parse_json('{"score": 0.95, "label": "relevant"}') == expected

    # Fenced JSON
    fenced = '```json\n{"score": 0.95, "label": "relevant"}\n```'
    assert parse_json(fenced) == expected

    # JSON surrounded by conversational chatter
    chatter = 'Here is the requested output:\n{"score": 0.95, "label": "relevant"}\nHope this helps!'
    assert parse_json(chatter) == expected

    # Garbage / malformed inputs
    assert parse_json("Invalid garbage text !!!") == {}
    assert parse_json("") == {}

    # JSON arrays (must return {})
    assert parse_json("[1, 2, 3]") == {}
    assert parse_json('[{"score": 0.95}]') == {}


def test_embed_and_caching(tmp_path: Path) -> None:
    """Verify embedding unit length, two-tier caching, and corrupt cache resilience."""
    client = install(tmp_path)
    texts = [
        "First document about gradient descent optimization.",
        "Second document regarding transformer attention layers.",
    ]

    # 1. First embed call: verify unit length and cache file written
    v1 = embed(texts, kind="document")
    assert v1.shape == (2, 128)
    norms = np.linalg.norm(v1, axis=1)
    assert np.allclose(norms, 1.0, atol=1e-5)

    # Identical text gives identical vectors
    v_single = embed([texts[0]], kind="document")
    assert np.allclose(v1[0], v_single[0], atol=1e-5)

    # Verify document cache written to disk
    cache_path = tmp_path / "embed_cache.json"
    assert cache_path.is_file()
    disk_data = json.loads(cache_path.read_text(encoding="utf-8"))
    assert len(disk_data) == 2

    # 2. Query embeddings are NOT written to disk
    call_count_before = len(client.calls)
    q_vec = embed(["Search query for transformers"], kind="query")
    assert q_vec.shape == (1, 128)
    assert np.allclose(np.linalg.norm(q_vec, axis=1), 1.0, atol=1e-5)
    # Disk cache still only has 2 document entries
    disk_data_after = json.loads(cache_path.read_text(encoding="utf-8"))
    assert len(disk_data_after) == 2

    # 3. Second call for identical texts makes zero new client calls
    calls_before = len(client.calls)
    v2 = embed(texts, kind="document")
    assert len(client.calls) == calls_before, "Cached document embed should make 0 client calls"
    assert np.allclose(v1, v2, atol=1e-5)

    # Second query call makes zero new client calls (in-memory cache)
    _ = embed(["Search query for transformers"], kind="query")
    assert len(client.calls) == calls_before, "Cached query embed should make 0 client calls"

    # 4. Corrupt cache file resilience
    reset_embedding_cache()
    cache_path.write_text("{ corrupt json data !!!", encoding="utf-8")
    v3 = embed(texts, kind="document")
    assert v3.shape == (2, 128)
    # File should have been safely overwritten with valid JSON
    new_data = json.loads(cache_path.read_text(encoding="utf-8"))
    assert len(new_data) == 2


def main() -> None:
    """Run all tests directly with plain asserts."""
    print("Running test_chunk_text...")
    test_chunk_text()

    with tempfile.TemporaryDirectory() as td:
        tpath = Path(td)
        print("Running test_markdown_ingest...")
        test_markdown_ingest(tpath)

    with tempfile.TemporaryDirectory() as td:
        tpath = Path(td)
        print("Running test_docx_and_pdf...")
        test_docx_and_pdf(tpath)

    print("Running test_load_chunks_sample_corpus...")
    test_load_chunks_sample_corpus()

    print("Running test_parse_json...")
    test_parse_json()

    with tempfile.TemporaryDirectory() as td:
        tpath = Path(td)
        print("Running test_embed_and_caching...")
        test_embed_and_caching(tpath)

    print("All ingest and llm tests passed successfully!")


if __name__ == "__main__":
    main()
